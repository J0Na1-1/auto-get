#!/usr/bin/env python3
"""auto-get: Daily security advisory crawler"""
import os, re, sys, time, logging
from datetime import datetime, timedelta, timezone
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

OUTPUT_DIR = '/CloudNAS/AI'
HK_UTC_OFFSET = 8
HK_TZ = timezone(timedelta(hours=HK_UTC_OFFSET))
NOW_HK = datetime.now(HK_TZ)
TODAY_HK = sys.argv[1] if len(sys.argv) > 1 else NOW_HK.strftime('%Y-%m-%d')
HEADERS = {'User-Agent': 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36'}

def normalize_platform(text):
    if not text: return 'Unknown'
    t = text.lower().strip()
    if re.match(r'forti', t): return 'Forti'
    if any(k in t for k in ['ios','macos','mac os','macintosh','iphone','ipad','ipod','apple']): return 'Apple'
    if 'microsoft' in t or re.search(r'windows\s*\d', t): return 'Microsoft'
    if 'cisco' in t: return 'Cisco'
    if any(k in t for k in ['palo alto','pan-os','palo']): return 'Palo Alto'
    if any(k in t for k in ['google','chrome','chromium']): return 'Google'
    if any(k in t for k in ['mozilla','firefox']): return 'Mozilla'
    if 'adobe' in t: return 'Adobe'
    if 'f5' in t: return 'F5'
    if 'red hat' in t: return 'Red Hat'
    if 'suse' in t: return 'SUSE'
    if 'ubuntu' in t: return 'Ubuntu'
    if 'debian' in t: return 'Debian'
    if 'oracle' in t: return 'Oracle'
    if 'vmware' in t: return 'VMware'
    if 'wordpress' in t: return 'WordPress'
    if t == 'linux': return 'Linux'
    return text.title()

def should_skip(risk_level):
    return risk_level in ['Low', 'Info', 'Informational']

def parse_date_hk(date_str):
    if not date_str: return None
    formats = ['%d-%B-%Y','%d %B %Y','%d %b %Y','%Y-%m-%d','%b %d, %Y','%d/%m/%Y','%m/%d/%Y','%d %b %Y %H:%M:%S','%Y-%m-%d %H:%M:%S','%d %b']
    for fmt in formats:
        try:
            dt = datetime.strptime(date_str.strip(), fmt)
            if dt.year == 1900:
                dt = dt.replace(year=NOW_HK.year)
            return dt.replace(tzinfo=timezone.utc).astimezone(HK_TZ)
        except ValueError: continue
    return None

def is_today_hk(dt_hk):
    if dt_hk is None: return False
    return dt_hk.strftime('%Y-%m-%d') == TODAY_HK

def crawl_govcert():
    results = []
    try:
        resp = requests.get('https://www.govcert.gov.hk/en/alerts.php', timeout=30, headers=HEADERS)
        soup = BeautifulSoup(resp.text, 'html.parser')
        alerts = soup.find_all('a', href=re.compile(r'alerts_detail\.php\?id=\d+'))
        for alert in alerts[:10]:
            title = alert.text.strip()
            href = alert.get('href', '')
            if href.startswith('http'):
                link = href
            elif href.startswith('/'):
                link = 'https://www.govcert.gov.hk' + href
            else:
                link = 'https://www.govcert.gov.hk/en/' + href
            parent = alert.parent
            date_text = parent.get_text() if parent else ''
            date_match = re.search(r'(\d{1,2}-\w+-\d{4}|\d{1,2}\s+\w+\s+\d{4})', date_text)
            pub_date = parse_date_hk(date_match.group(1)) if date_match else None
            if not is_today_hk(pub_date): continue
            detail = fetch_detail_govcert(link, title)
            if detail: results.append(detail)
    except Exception as e:
        logger.error(f'GovCERT crawl error: {e}')
    return results

def fetch_detail_govcert(url, title):
    try:
        resp = requests.get(url, timeout=30, headers=HEADERS)
        soup = BeautifulSoup(resp.text, 'html.parser')
        text = soup.get_text()
        dm = re.search(r'Published on:\s*(.+)', text)
        pub_date = dm.group(1).strip() if dm else ''
        descm = re.search(r'Description:\s*(.+?)(?=Affected Systems|Impact|Recommendation|More Information|\Z)', text, re.DOTALL)
        description = descm.group(1).strip()[:500] if descm else ''
        affm = re.search(r'Affected Systems:\s*(.+?)(?=Impact|Recommendation|More Information|\Z)', text, re.DOTALL)
        affected = affm.group(1).strip() if affm else ''
        recm = re.search(r'Recommendation:\s*(.+?)(?=More Information|\Z)', text, re.DOTALL)
        recommendation = recm.group(1).strip()[:500] if recm else ''
        cve_links = soup.find_all('a', href=re.compile(r'cve\.mitre\.org'))
        cve = '; '.join([a.text.strip() for a in cve_links])
        tags = soup.find_all('a', href=re.compile(r'/en/alerts\.php\?tag='))
        skip_tags = ['PHP','Edge','Google','Chrome','Cisco','Firefox','Mozilla','Windows','Windows Server','Microsoft Office','Apple','iPod','Visual Studio','VMware','Internet Explorer','Microsoft 365 Apps','Android','F5','Acrobat','Acrobat Reader','Adobe Reader','Word','Previous']
        platform_tags = [a.text.strip() for a in tags if a.text.strip() not in skip_tags]
        platform = '; '.join(platform_tags) if platform_tags else affected
        return {'Date': pub_date, 'Risk level': 'Medium',
            'Platform': normalize_platform(platform) if platform else 'Unknown',
            'Affected Product': affected[:500] if affected else '',
            'Description': (title + ' - ' + description) if description else title,
            'Workaround': recommendation[:500] if recommendation else '',
            'Related Link': url, 'CVE': cve, 'From': 'GovCERT'}
    except Exception as e:
        logger.error(f'GovCERT detail error: {e}')
        return None

def crawl_hkcert():
    results = []
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, executable_path='/usr/bin/chromium')
            page = browser.new_page()
            page.goto('https://www.hkcert.org/security-bulletin/', timeout=30000)
            time.sleep(5)
            html = page.content()
            links = re.findall(r'href=\"(/security-bulletin/[^\"]+)\"', html)
            links = [l for l in links if not l.endswith('#main') and not l.startswith('/tc/')]
            for link in links[:15]:
                try:
                    full_link = link if link.startswith('http') else 'https://www.hkcert.org' + link
                    idx = html.find(link)
                    if idx >= 0:
                        surrounding = html[max(0, idx-300):idx+300]
                        date_match = re.search(r'(Release Date|Last Update Date):\s*(\d{1,2}\s+\w+\s+\d{4})', surrounding)
                        pub_date = parse_date_hk(date_match.group(2)) if date_match else None
                    else:
                        pub_date = None
                    if not is_today_hk(pub_date): continue
                    page.goto(full_link, timeout=30000)
                    time.sleep(3)
                    detail_html = page.content()
                    detail_soup = BeautifulSoup(detail_html, 'html.parser')
                    detail_text = detail_soup.get_text()
                    title = link.split('/')[-1].replace('_', ' ').title()
                    descm = re.search(r'Description\s*[:\n]\s*(.+?)(?=Impact|System|Solution|Vulnerability Identifier|Related Links|\Z)', detail_text, re.DOTALL)
                    description = descm.group(1).strip()[:500] if descm else ''
                    affm = re.search(r'System / Technologies affected\s*[:\n]\s*(.+?)(?=Impact|Solutions|Vulnerability Identifier|Related Links|\Z)', detail_text, re.DOTALL)
                    affected = affm.group(1).strip() if affm else ''
                    solm = re.search(r'Solutions?\s*[:\n]\s*(.+?)(?=Vulnerability Identifier|Related Links|\Z)', detail_text, re.DOTALL)
                    workaround = solm.group(1).strip()[:500] if solm else ''
                    cve_match = re.findall(r'CVE-\d{4}-\d{4,7}', detail_text)
                    cve = '; '.join(cve_match)
                    results.append({'Date': pub_date.strftime('%d %B %Y') if pub_date else '',
                        'Risk level': 'Medium',
                        'Platform': normalize_platform(affected) if affected else 'Unknown',
                        'Affected Product': affected[:500] if affected else '',
                        'Description': description if description else title,
                        'Workaround': workaround if workaround else '',
                        'Related Link': full_link, 'CVE': cve, 'From': 'HKCERT'})
                except Exception as e:
                    logger.error(f'HKCERT detail error: {e}')
                    continue
            browser.close()
    except Exception as e:
        logger.error(f'HKCERT crawl error: {e}')
    return results

def crawl_fortinet():
    results = []
    try:
        resp = requests.get('https://www.fortiguard.com/psirt', timeout=30, headers=HEADERS)
        soup = BeautifulSoup(resp.text, 'html.parser')
        text = soup.get_text()
        advisories = re.findall(r'(FG-IR-26-\d+)\s+(CVE-\d{4}-\d+)', text)
        for adv_id, cve in advisories[:20]:
            pattern = rf'{re.escape(adv_id)}.*?{re.escape(cve)}(.*?)(?=FG-IR-26-|Published:|\Z)'
            match = re.search(pattern, text, re.DOTALL)
            context = match.group(1) if match else ''
            title_match = re.search(r'(FG-IR-26-\d+.*?)(?:CVE-|Published:)', context)
            title = title_match.group(1).strip() if title_match else adv_id
            sev_match = re.search(r'(Critical|High|Medium|Low)\s+Severity', context, re.I)
            severity = sev_match.group(1) if sev_match else 'Medium'
            date_match = re.search(r'Published:\s*(Sep \d+, \d+|Jul \d+, \d+|Aug \d+, \d+|Oct \d+, \d+|Jun \d+, \d+|May \d+, \d+|Apr \d+, \d+|Mar \d+, \d+|Feb \d+, \d+|Jan \d+, \d+|Nov \d+, \d+|Dec \d+, \d+)', context)
            pub_date = parse_date_hk(date_match.group(1)) if date_match else None
            if not is_today_hk(pub_date): continue
            if should_skip(severity): continue
            desc_match = re.search(r'(?:Improper|NULL|Uncontrolled|Use of|Path traversal|XXE|SQL Injection|Command Injection|Buffer Overflow|Out-of-bounds|Deserialization|Race Condition|Resource Consumption|Cross-site Scripting|HTTP/2|TLS|Session|Certificate|Authorization|Authentication|Memory|Integer|Overflow|Injection|Bypass|Exposure|Misconfiguration)\s+(.+?)(?:CVE-|Published:|Severity)', context, re.DOTALL)
            description = desc_match.group(1).strip()[:300] if desc_match else title
            prod_match = re.search(r'(FortiOS|FortiProxy|FortiAnalyzer|FortiManager|FortiSwitch|FortiAP|FortiWeb|FortiSIEM|FortiMail|FortiPortal|FortiSandbox|FortiClient|FortiAuthenticator|FortiDDoS|FortiExtender|FortiDeceiver|FortiNAC|FortiWLC|FortiConverter|FortiPresence|FortiVoice|FortiTester|FortiSOAR|FortiPAM|FortiProxy|FortiSwitchManager|FortiGate)\s[\d\.\s,]+', context)
            affected = prod_match.group(0).strip() if prod_match else ''
            results.append({'Date': date_match.group(1) if date_match else '',
                'Risk level': severity, 'Platform': 'Forti',
                'Affected Product': affected[:500] if affected else '',
                'Description': description,
                'Workaround': 'Apply vendor patches from FortiGuard PSIRT',
                'Related Link': 'https://www.fortiguard.com/psirt',
                'CVE': cve, 'From': 'Fortinet'})
    except Exception as e:
        logger.error(f'Fortinet crawl error: {e}')
    return results

def crawl_paloalto():
    results = []
    try:
        csv_url = 'https://security.paloaltonetworks.com/csv'
        resp = requests.get(csv_url, timeout=30, headers=HEADERS)
        if resp.status_code == 200:
            lines = resp.text.strip().split('\n')
            if len(lines) > 1:
                for line in lines[1:]:
                    fields = line.split(',')
                    if len(fields) >= 7:
                        published = fields[5].strip() if len(fields) > 5 else ''
                        pub_date = parse_date_hk(published)
                        if not is_today_hk(pub_date): continue
                        cve = fields[0].strip() if fields[0] else ''
                        severity = 'Medium'
                        cvss_match = re.search(r'(\d+\.?\d*)', fields[0] if fields else '')
                        if cvss_match:
                            cvss = float(cvss_match.group(1))
                            if cvss >= 9.0: severity = 'Critical'
                            elif cvss >= 7.0: severity = 'High'
                            elif cvss >= 4.0: severity = 'Medium'
                            else: severity = 'Low'
                        if should_skip(severity): continue
                        results.append({'Date': published, 'Risk level': severity,
                            'Platform': 'Palo Alto',
                            'Affected Product': fields[2].strip()[:200] if len(fields) > 2 else '',
                            'Description': fields[1].strip()[:500] if len(fields) > 1 else '',
                            'Workaround': 'Apply vendor patches',
                            'Related Link': f'https://security.paloaltonetworks.com/{cve}',
                            'CVE': cve, 'From': 'Palo Alto'})
    except Exception as e:
        logger.error(f'Palo Alto crawl error: {e}')
    return results

def crawl_cisco():
    results = []
    try:
        url = 'https://sec.cloudapps.cisco.com/security/center/publicationService.x?criteria=exact&cves=&keyword=&last_published_date=&limit=50&offset=0&publicationTypeIDs=1,3&securityImpactRatings=&sort=-day_sir&title='
        resp = requests.get(url, timeout=30, headers=HEADERS)
        if resp.status_code == 200 and len(resp.text) > 0:
            data = resp.json()
            for adv in data[:20]:
                published = adv.get('lastPublished', '')
                pub_date = parse_date_hk(published[:10] if published else '') if published else None
                if not is_today_hk(pub_date): continue
                severity = adv.get('severity', 'Medium')
                if severity.lower() == 'critical': severity = 'Critical'
                elif severity.lower() == 'high': severity = 'High'
                elif severity.lower() == 'medium': severity = 'Medium'
                elif severity.lower() == 'low': severity = 'Low'
                else: severity = 'Medium'
                if should_skip(severity): continue
                cve = adv.get('cve', '')
                title = adv.get('title', '')
                results.append({'Date': pub_date.strftime('%d %B %Y') if pub_date else (published[:10] if published else ''),
                    'Risk level': severity,
                    'Platform': 'Cisco',
                    'Affected Product': title[:200],
                    'Description': title,
                    'Workaround': 'Apply vendor patches from Cisco PSIRT',
                    'Related Link': adv.get('url', ''),
                    'CVE': cve, 'From': 'Cisco'})
    except Exception as e:
        logger.error(f'Cisco crawl error: {e}')
    return results

def write_xlsx(results):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Advisories'
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0ea5e9', end_color='0ea5e9', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    headers = ['Date', 'Risk level', 'Platform', 'Affected Product', 'Description', 'Workaround', 'Related Link', 'CVE', 'From']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    for row_idx, item in enumerate(results, 2):
        for col_idx, key in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=item.get(key, ''))
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    widths = [12, 14, 30, 50, 45, 40, 30, 12, 12]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f'AlertNews{datetime.strptime(TODAY_HK, "%Y-%m-%d").strftime("%d%m%y")}.xlsx'
    filepath = os.path.join(OUTPUT_DIR, filename)
    wb.save(filepath)
    logger.info(f'Output saved to {filepath}')
    return filepath

def main():
    logger.info(f'=== auto-get started at {NOW_HK.strftime("%Y-%m-%d %H:%M:%S")} HK Time ===')
    all_results = []
    sites = [('GovCERT', crawl_govcert), ('HKCERT', crawl_hkcert), ('Fortinet', crawl_fortinet), ('Palo Alto', crawl_paloalto), ('Cisco', crawl_cisco)]
    for name, crawler in sites:
        logger.info(f'Crawling {name}...')
        try:
            results = crawler()
            all_results.extend(results)
            logger.info(f'  {name}: {len(results)} new advisories')
        except Exception as e:
            logger.error(f'  {name} error: {e}')
    all_results.sort(key=lambda x: x.get('Date', ''), reverse=True)
    logger.info(f'Total: {len(all_results)} advisories for {TODAY_HK}')
    filepath = write_xlsx(all_results)
    logger.info(f'=== auto-get completed ===')
    return filepath

if __name__ == '__main__':
    main()